"""Export knihy jízd do XLSX, CSV a PDF (zadání 23).

Sloupce jsou definované **jednou** (`COLUMNS`) a všechny tři formáty z
nich vycházejí. Kdyby si každý držel vlastní seznam, po prvním přidaném
údaji by se rozešly a nikdo by si toho chvíli nevšiml.

Řádek se staví čistou funkcí `row_for(trip)` nad už načtenou jízdou -
žádná databáze, takže se dá testovat přímo a export nikdy nevygeneruje
jiná data, než jaká byla ve výpisu.

Hodnoty se drží jako čísla, ne jako formátované řetězce: v XLSX se s
nimi pak dá počítat. Na text se převádějí až v CSV a PDF, kde nic jiného
neexistuje.
"""
import csv
import io
from dataclasses import dataclass
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.lib.styles import ParagraphStyle

from app.core import labels
from app.core.pdf_font import fold, get_pdf_font

LOCAL_TZ = ZoneInfo("Europe/Prague")


@dataclass(frozen=True)
class Column:
    key: str
    header: str
    #: Jen tyhle se vejdou do PDF - viz `to_pdf`.
    in_pdf: bool = False
    #: Šířka sloupce v XLSX (znaky) a v PDF (mm).
    xlsx_width: int = 16
    pdf_width: float = 20.0


# Pořadí a rozsah podle zadání 22 („záznam obsahuje minimálně...").
COLUMNS = [
    Column("date", "Datum", in_pdf=True, xlsx_width=11, pdf_width=17),
    Column("vehicle", "Vozidlo", in_pdf=True, xlsx_width=14, pdf_width=22),
    Column("driver", "Řidič", in_pdf=True, xlsx_width=20, pdf_width=30),
    Column("extra_drivers", "Další řidiči", xlsx_width=22),
    Column("started_at", "Začátek", in_pdf=True, xlsx_width=16, pdf_width=17),
    Column("ended_at", "Konec", in_pdf=True, xlsx_width=16, pdf_width=17),
    Column("start_km", "Počáteční km", xlsx_width=13),
    Column("end_km", "Konečné km", xlsx_width=13),
    Column("driven_km", "Ujeté km", in_pdf=True, xlsx_width=10, pdf_width=17),
    Column("start_fuel", "Nádrž na začátku (%)", xlsx_width=19),
    Column("end_fuel", "Nádrž na konci (%)", xlsx_width=19),
    Column("purpose", "Účel", in_pdf=True, xlsx_width=18, pdf_width=26),
    Column("route", "Trasa", in_pdf=True, xlsx_width=30, pdf_width=48),
    Column("fuelings", "Tankování / nabíjení", xlsx_width=30),
    Column("fuelings_amount", "Za palivo (Kč)", xlsx_width=14),
    Column("status", "Stav", xlsx_width=11),
    Column("notes", "Poznámky", xlsx_width=34),
    Column("defects", "Závady", xlsx_width=28),
]


def _local(value: datetime | None) -> datetime | None:
    return value.astimezone(LOCAL_TZ) if value is not None else None


def _format_dt(value: datetime | None) -> str:
    local = _local(value)
    return f"{local:%d.%m.%Y %H:%M}" if local else ""


def row_for(trip) -> dict:
    """Jeden řádek knihy jízd. Čistá funkce nad načtenou jízdou."""
    started = _local(trip.started_at)
    driven = (
        trip.end_odometer_km - trip.start_odometer_km
        if trip.end_odometer_km is not None else None
    )

    fuelings = [
        " ".join(part for part in (
            f"{float(f.quantity):.2f}".rstrip("0").rstrip(".").replace(".", ","),
            f.unit,
            f"({f.station})" if f.station else "",
        ) if part)
        for f in trip.fuelings
    ]
    fuel_amount = sum(float(f.price_total_czk) for f in trip.fuelings if f.price_total_czk)

    purpose = labels.TRIP_PURPOSE.get(trip.purpose_code, trip.purpose_code or "")
    if trip.purpose_text:
        purpose = f"{purpose} – {trip.purpose_text}" if purpose else trip.purpose_text

    # Poznámka řidiče při uzavření a pozdější komentáře jsou dvě různé
    # věci (viz model TripNote), ale ve výpisu je zajímá obojí.
    notes = [trip.note] if trip.note else []
    notes += [note.text for note in trip.notes]

    return {
        "date": f"{started:%d.%m.%Y}" if started else "",
        "vehicle": f"{trip.vehicle.license_plate} ({trip.vehicle.brand} {trip.vehicle.model})".strip(),
        "driver": trip.primary_driver.full_name if trip.primary_driver else "",
        "extra_drivers": ", ".join(
            d.user.full_name for d in trip.extra_drivers if d.user
        ),
        "started_at": _format_dt(trip.started_at),
        "ended_at": _format_dt(trip.ended_at),
        "start_km": trip.start_odometer_km,
        "end_km": trip.end_odometer_km,
        "driven_km": driven,
        "start_fuel": trip.start_fuel_level,
        "end_fuel": trip.end_fuel_level,
        "purpose": purpose,
        "route": trip.route_text or "",
        "fuelings": "; ".join(fuelings),
        "fuelings_amount": round(fuel_amount, 2) if fuel_amount else None,
        "status": labels.TRIP_STATUS.get(trip.status, trip.status),
        "notes": " | ".join(notes),
        # VehicleDefect nemá název, jen popis - ten je to, co člověk
        # v exportu hledá ("prasklé zrcátko"), i s prioritou.
        "defects": "; ".join(
            f"{defect.description} ({labels.DEFECT_PRIORITY.get(defect.priority, defect.priority)})"
            for defect in trip.defects
        ),
    }


def build_rows(trips) -> list[dict]:
    return [row_for(trip) for trip in trips]


def _as_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:.2f}".replace(".", ",")
    return str(value)


# --- CSV ---------------------------------------------------------------

def to_csv(rows: list[dict]) -> bytes:
    """Středník a BOM, ne čárka a holé UTF-8.

    Český Excel otevře CSV s čárkou jako jeden sloupec a bez BOM zobrazí
    diakritiku rozsypanou. Export, který se musí před použitím opravovat,
    je k ničemu - a kdo ho čte skriptem, si se středníkem poradí."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
    writer.writerow([column.header for column in COLUMNS])
    for row in rows:
        writer.writerow([_as_text(row.get(column.key)) for column in COLUMNS])
    return b"\xef\xbb\xbf" + buffer.getvalue().encode("utf-8")


# --- XLSX --------------------------------------------------------------

def to_xlsx(rows: list[dict], *, title: str = "Kniha jízd") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    # Název listu nesmí mít víc než 31 znaků ani lomítka.
    sheet.title = title[:31].replace("/", "-")

    sheet.append([column.header for column in COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for row in rows:
        # Čísla jako čísla, ne jako text - jinak v Excelu nejde sečíst
        # sloupec ujetých km, což je první věc, kterou tam kdo udělá.
        sheet.append([row.get(column.key) for column in COLUMNS])

    for index, column in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = column.xlsx_width

    # Záhlaví zůstane viditelné při rolování; bez toho se u padesáti
    # jízd nedá poznat, který sloupec je který.
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --- PDF ---------------------------------------------------------------

PDF_COLUMNS = [column for column in COLUMNS if column.in_pdf]


def to_pdf(rows: list[dict], *, title: str = "Kniha jízd", subtitle: str = "") -> bytes:
    """PDF je na čtení, ne na další zpracování.

    Proto se do něj vysází jen část sloupců (`in_pdf`): osmnáct sloupců
    na šířku A4 by dalo písmo, které nikdo nepřečte. Kdo potřebuje
    všechno, sáhne po XLSX nebo CSV."""
    font = get_pdf_font()

    def text(value) -> str:
        return fold(_as_text(value), font)

    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
        title=text(title), author="Kniha jízd",
    )

    heading = ParagraphStyle("kj-heading", fontName=font.bold, fontSize=15, leading=19)
    sub = ParagraphStyle("kj-sub", fontName=font.regular, fontSize=9, leading=12,
                         textColor=colors.HexColor("#64748B"))
    cell = ParagraphStyle("kj-cell", fontName=font.regular, fontSize=7.5, leading=9.5)
    head = ParagraphStyle("kj-head", fontName=font.bold, fontSize=7.5, leading=9.5,
                          textColor=colors.white)

    story = [Paragraph(text(title), heading)]
    if subtitle:
        story.append(Paragraph(text(subtitle), sub))
    story.append(Spacer(1, 6 * mm))

    # Paragraph v buňkách, ne holý řetězec: dlouhá trasa se musí zalomit,
    # jinak přeteče přes sousední sloupec.
    data = [[Paragraph(text(column.header), head) for column in PDF_COLUMNS]]
    for row in rows:
        data.append([Paragraph(text(row.get(column.key)), cell) for column in PDF_COLUMNS])

    table = Table(
        data,
        colWidths=[column.pdf_width * mm for column in PDF_COLUMNS],
        repeatRows=1,  # záhlaví se zopakuje na každé stránce
    )
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0E7490")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.append(table)

    if not rows:
        story.append(Spacer(1, 6 * mm))
        story.append(Paragraph(text("Zadanému filtru neodpovídá žádná jízda."), sub))

    document.build(story)
    return buffer.getvalue()
