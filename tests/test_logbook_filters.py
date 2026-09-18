"""Filtr a exportní řádky - čisté funkce bez databáze (zadání 22/23)."""
import io
import uuid
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

from openpyxl import load_workbook

from app.modules.logbook import export
from app.modules.logbook.filters import LOCAL_TZ, LogbookFilter


# --- parsování filtru --------------------------------------------------

def test_empty_params_mean_no_filter():
    flt = LogbookFilter.from_params({})
    assert flt.is_empty
    assert flt.page == 1
    assert flt.started_from is None and flt.started_to is None


def test_period_bounds_cover_the_whole_last_day():
    flt = LogbookFilter.from_params({"from": "2026-09-01", "to": "2026-09-30"})
    assert flt.started_from == datetime(2026, 9, 1, tzinfo=LOCAL_TZ)
    # Horní mez je půlnoc PO posledním dni - jinak by jízda z 30. 9. v
    # 17:00 vypadla.
    assert flt.started_to == datetime(2026, 10, 1, tzinfo=LOCAL_TZ)


def test_reversed_period_is_swapped_not_rejected():
    flt = LogbookFilter.from_params({"from": "2026-09-30", "to": "2026-09-01"})
    assert flt.date_from == date(2026, 9, 1)
    assert flt.date_to == date(2026, 9, 30)


def test_garbage_is_ignored_instead_of_crashing():
    """Rozbitý parametr v URL nemá být konec světa - filtr se nepoužije."""
    flt = LogbookFilter.from_params({
        "from": "vcera", "to": "", "vehicle": "neni-uuid",
        "driver": "12345", "purpose": "vymysleny", "status": "smazana", "page": "-3",
    })
    assert flt.is_empty
    assert flt.page == 1


def test_known_codes_survive():
    vehicle_id, driver_id = uuid.uuid4(), uuid.uuid4()
    flt = LogbookFilter.from_params({
        "vehicle": str(vehicle_id), "driver": str(driver_id),
        "purpose": "montaz", "status": "completed", "page": "3",
    })
    assert flt.vehicle_id == vehicle_id
    assert flt.driver_id == driver_id
    assert flt.purpose_code == "montaz"
    assert flt.status == "completed"
    assert flt.page == 3 and flt.offset == 100
    assert not flt.is_empty


def test_query_string_keeps_the_filter_and_drops_defaults():
    flt = LogbookFilter.from_params({"from": "2026-09-01", "purpose": "servis"})
    query = flt.query_string()
    assert "from=2026-09-01" in query
    assert "purpose=servis" in query
    # Prázdné hodnoty ani "page=1" v URL nemají co dělat.
    assert "to=" not in query and "page=" not in query

    assert "page=4" in flt.query_string(page=4)
    assert "page" not in flt.query_string(page=None)


# --- exportní řádek ----------------------------------------------------

@dataclass
class FakeUser:
    full_name: str


@dataclass
class FakeVehicle:
    license_plate: str = "1AB 2345"
    brand: str = "Škoda"
    model: str = "Octavia"
    fuel_type: str = "nafta"


@dataclass
class FakeFueling:
    quantity: float
    unit: str = "l"
    price_total_czk: float | None = None
    station: str | None = None


@dataclass
class FakeExtraDriver:
    user: FakeUser | None


@dataclass
class FakeNote:
    text: str


@dataclass
class FakeDefect:
    description: str
    priority: str = "normal"


@dataclass
class FakeTrip:
    vehicle: FakeVehicle = field(default_factory=FakeVehicle)
    primary_driver: FakeUser | None = field(default_factory=lambda: FakeUser("Karel Novák"))
    status: str = "completed"
    started_at: datetime = datetime(2026, 9, 14, 8, 0, tzinfo=LOCAL_TZ)
    ended_at: datetime | None = datetime(2026, 9, 14, 16, 30, tzinfo=LOCAL_TZ)
    start_odometer_km: int = 100000
    end_odometer_km: int | None = 100120
    start_fuel_level: int | None = 80
    end_fuel_level: int | None = 45
    purpose_code: str | None = "montaz"
    purpose_text: str | None = None
    route_text: str | None = "Mladá Boleslav – Zlatá Olešnice"
    note: str | None = None
    extra_drivers: list = field(default_factory=list)
    fuelings: list = field(default_factory=list)
    notes: list = field(default_factory=list)
    defects: list = field(default_factory=list)


def test_row_has_every_field_the_spec_asks_for():
    trip = FakeTrip(
        extra_drivers=[FakeExtraDriver(FakeUser("Petr Dvořák"))],
        fuelings=[FakeFueling(42.5, "l", 1275.0, "ONO Mladá Boleslav")],
        note="Zácpa na D10",
        notes=[FakeNote("Doplněno odpoledne")],
        defects=[FakeDefect("Prasklé zpětné zrcátko", "high")],
    )
    row = export.row_for(trip)

    assert row["date"] == "14.09.2026"
    assert row["vehicle"] == "1AB 2345 (Škoda Octavia)"
    assert row["driver"] == "Karel Novák"
    assert row["extra_drivers"] == "Petr Dvořák"
    assert row["started_at"] == "14.09.2026 08:00"
    assert row["ended_at"] == "14.09.2026 16:30"
    assert row["driven_km"] == 120
    assert row["start_fuel"] == 80 and row["end_fuel"] == 45
    assert row["purpose"] == "Montáž"
    assert row["route"] == "Mladá Boleslav – Zlatá Olešnice"
    assert "42,5 l" in row["fuelings"] and "ONO Mladá Boleslav" in row["fuelings"]
    assert row["fuelings_amount"] == 1275.0
    assert row["status"] == "Ukončená"
    # Poznámka řidiče i pozdější komentář - dvě různé věci, obojí zajímá.
    assert "Zácpa na D10" in row["notes"] and "Doplněno odpoledne" in row["notes"]
    assert row["defects"] == "Prasklé zpětné zrcátko (Vysoká)"


def test_running_trip_has_no_driven_km():
    """U probíhající jízdy není konečný tachometr - nula by lhala."""
    row = export.row_for(FakeTrip(ended_at=None, end_odometer_km=None, status="active"))
    assert row["driven_km"] is None
    assert row["ended_at"] == ""
    assert row["status"] == "Probíhá"


def test_electric_trip_keeps_kwh():
    row = export.row_for(FakeTrip(fuelings=[FakeFueling(38.0, "kWh", 342.0)]))
    assert "38 kWh" in row["fuelings"]


def test_custom_purpose_text_is_appended():
    row = export.row_for(FakeTrip(purpose_code="jine", purpose_text="Odvoz na letiště"))
    assert row["purpose"] == "Jiný – Odvoz na letiště"


# --- formáty -----------------------------------------------------------

def test_csv_has_bom_and_semicolons():
    """Český Excel jinak otevře soubor jako jeden sloupec s rozsypanou
    diakritikou."""
    data = export.to_csv(export.build_rows([FakeTrip()]))
    assert data.startswith(b"\xef\xbb\xbf")

    text = data.decode("utf-8-sig")
    header, first = text.splitlines()[0], text.splitlines()[1]
    assert header.startswith("Datum;Vozidlo;Řidič")
    assert "1AB 2345" in first


def test_csv_column_count_matches_the_header():
    rows = export.build_rows([FakeTrip(), FakeTrip(ended_at=None, end_odometer_km=None)])
    lines = export.to_csv(rows).decode("utf-8-sig").strip().splitlines()
    assert len(lines) == 3
    assert all(line.count(";") == lines[0].count(";") for line in lines)


def test_xlsx_keeps_numbers_as_numbers():
    """Aby v Excelu šel sloupec ujetých km sečíst."""
    data = export.to_xlsx(export.build_rows([FakeTrip()]))
    sheet = load_workbook(io.BytesIO(data)).active

    headers = [cell.value for cell in sheet[1]]
    assert headers == [column.header for column in export.COLUMNS]

    driven = sheet.cell(row=2, column=headers.index("Ujeté km") + 1).value
    assert driven == 120
    assert isinstance(driven, int)
    # Záhlaví zůstane při rolování viditelné.
    assert sheet.freeze_panes == "A2"


def test_pdf_is_a_pdf_and_survives_an_empty_result():
    header = export.to_pdf([])[:5]
    assert header == b"%PDF-"

    with_rows = export.to_pdf(export.build_rows([FakeTrip()]), subtitle="celé období · 1 jízd")
    assert with_rows.startswith(b"%PDF-")
    # Víc dat = větší soubor; kdyby se řádky nevysázely, bylo by to stejné.
    assert len(with_rows) > len(export.to_pdf([]))


def test_all_three_formats_use_the_same_columns():
    """Sloupce jsou definované jednou - PDF jich bere podmnožinu."""
    assert export.PDF_COLUMNS, "PDF musí mít aspoň jeden sloupec"
    assert all(column in export.COLUMNS for column in export.PDF_COLUMNS)
    assert len(export.PDF_COLUMNS) < len(export.COLUMNS)
