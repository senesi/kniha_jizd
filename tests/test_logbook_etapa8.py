"""Etapa 8 - kniha jízd, filtry a exporty (zadání 22/23).

Těžiště je na dvou věcech: že filtr platí i pro export a že se skryté
vozidlo (požadavek D) nedostane ani do výpisu, ani do součtů, ani do
staženého souboru.
"""
import io
import uuid
from datetime import date, datetime, timedelta, timezone

from openpyxl import load_workbook
from sqlalchemy import select

from tests.conftest import create_vehicle, extract_csrf_token, login


async def _start(ac, csrf, vehicle_id, **fields):
    data = {"csrf_token": csrf, "start_odometer_km": "100000", "start_fuel_level": "50", **fields}
    response = await ac.post(f"/kniha-jizd/vehicles/{vehicle_id}/trips/start", data=data, follow_redirects=False)
    assert response.status_code == 303, response.text
    return response.headers["location"].split("/trips/")[1].split("?")[0]


def _has_trip(html: str, trip_id: str) -> bool:
    """Hledá odkaz na jízdu, ne SPZ.

    SPZ je na stránce vždycky - je i v rozbalovacím výběru vozidel, takže
    „SPZ v textu" by u filtru neznamenalo vůbec nic."""
    return f'/kniha-jizd/trips/{trip_id}"' in html


async def _end(ac, csrf, trip_id, **fields):
    data = {
        "csrf_token": csrf, "end_odometer_km": "100120", "end_fuel_level": "30",
        "purpose_code": "montaz", "route_text": "Mladá Boleslav – Zlatá Olešnice",
        **fields,
    }
    response = await ac.post(f"/kniha-jizd/trips/{trip_id}/end", data=data, follow_redirects=False)
    assert response.status_code == 303, response.text
    return response


async def _completed_trip(ac, csrf, vehicle_id, *, start_km=100000, driven=120, **end_fields) -> str:
    """Jízda musí navázat na stav tachometru předchozí - jinak ji
    aplikace odmítne (a má pravdu)."""
    trip_id = await _start(ac, csrf, vehicle_id, start_odometer_km=str(start_km))
    await _end(ac, csrf, trip_id, end_odometer_km=str(start_km + driven), **end_fields)
    return trip_id


async def _set_started_at(trip_id: str, moment: datetime) -> None:
    """Posune začátek jízdy do minulosti.

    Formulář to neumí (a schválně - jízda začíná teď), ale filtr podle
    období se jinak nedá vyzkoušet na jednom běhu testů."""
    from app.core.db import async_session_factory
    from app.models.fleet import Trip

    async with async_session_factory() as db:
        trip = (await db.execute(select(Trip).where(Trip.id == uuid.UUID(trip_id)))).scalar_one()
        trip.started_at = moment
        await db.commit()


def _csv_lines(content: bytes) -> list[str]:
    return content.decode("utf-8-sig").strip().splitlines()


# --- přístup -----------------------------------------------------------

async def test_driver_has_no_access_to_the_logbook(anon_client, basic_user):
    """Kniha jízd napříč vozidly je za fleet.logbook.view, ne pro každého."""
    await login(anon_client, basic_user)
    assert (await anon_client.get("/kniha-jizd/logbook")).status_code == 403
    for extension in ("csv", "xlsx", "pdf"):
        assert (await anon_client.get(f"/kniha-jizd/logbook/export.{extension}")).status_code == 403


async def test_responsible_person_has_access(anon_client, responsible_user):
    await login(anon_client, responsible_user)
    assert (await anon_client.get("/kniha-jizd/logbook")).status_code == 200


async def test_nav_shows_the_logbook_only_to_those_who_may_see_it(
    logged_in_client, anon_client, basic_user,
):
    admin_page = await logged_in_client.get("/kniha-jizd/")
    assert "/kniha-jizd/logbook" in admin_page.text

    await login(anon_client, basic_user)
    driver_page = await anon_client.get("/kniha-jizd/")
    assert "/kniha-jizd/logbook" not in driver_page.text


async def test_unknown_export_format_is_404(logged_in_client):
    assert (await logged_in_client.get("/kniha-jizd/logbook/export.exe")).status_code == 404


# --- výpis a souhrn ----------------------------------------------------

async def test_logbook_lists_trips_with_summary(logged_in_client, csrf_token):
    vehicle_id = await create_vehicle(
        logged_in_client, csrf_token, internal_code="L-01", license_plate="1LB 0001",
    )
    trip_id = await _completed_trip(logged_in_client, csrf_token, vehicle_id)

    page = await logged_in_client.get(f"/kniha-jizd/logbook?vehicle={vehicle_id}")
    assert page.status_code == 200
    assert _has_trip(page.text, trip_id)
    assert "1LB 0001" in page.text
    assert "Mladá Boleslav" in page.text
    assert "120" in page.text


async def test_running_trip_is_not_counted_into_driven_km(logged_in_client, csrf_token):
    """U probíhající jízdy není konečný tachometr - nula by součet snížila."""
    from app.modules.logbook import repository
    from app.modules.logbook.filters import LogbookFilter
    from app.core.db import async_session_factory
    from app.core.deps import get_user_permission_codes
    from app.models.core import User

    vehicle_id = await create_vehicle(
        logged_in_client, csrf_token, internal_code="L-02", license_plate="1LB 0002",
    )
    await _start(logged_in_client, csrf_token, vehicle_id)

    async with async_session_factory() as db:
        admin = (await db.execute(select(User).order_by(User.created_at))).scalars().first()
        codes = await get_user_permission_codes(db, admin.id)
        flt = LogbookFilter(vehicle_id=uuid.UUID(vehicle_id))
        result = await repository.summary(db, flt, codes=codes, user=admin)

    assert result["count"] == 1
    assert result["active"] == 1
    assert result["total_km"] == 0


# --- filtry ------------------------------------------------------------

async def test_period_filter_limits_both_list_and_export(logged_in_client, csrf_token):
    vehicle_id = await create_vehicle(
        logged_in_client, csrf_token, internal_code="L-10", license_plate="1LB 0010",
    )
    old_trip = await _completed_trip(logged_in_client, csrf_token, vehicle_id)
    await _set_started_at(old_trip, datetime(2026, 1, 15, 9, 0, tzinfo=timezone.utc))
    new_trip = await _completed_trip(
        logged_in_client, csrf_token, vehicle_id,
        start_km=100120, driven=180, route_text="Praha – Brno",
    )

    narrow = f"/kniha-jizd/logbook?vehicle={vehicle_id}&from=2026-01-01&to=2026-01-31"

    all_rows = _csv_lines((await logged_in_client.get(f"/kniha-jizd/logbook/export.csv?vehicle={vehicle_id}")).content)
    narrow_rows = _csv_lines((await logged_in_client.get(
        f"/kniha-jizd/logbook/export.csv?vehicle={vehicle_id}&from=2026-01-01&to=2026-01-31"
    )).content)

    assert len(all_rows) == 3      # hlavička + dvě jízdy
    assert len(narrow_rows) == 2   # hlavička + ta lednová
    assert "15.01.2026" in narrow_rows[1]
    assert "Praha – Brno" not in "\n".join(narrow_rows)

    page = await logged_in_client.get(narrow)
    assert _has_trip(page.text, old_trip)
    assert not _has_trip(page.text, new_trip)


async def test_driver_filter_finds_extra_drivers_too(logged_in_client, csrf_token, basic_user):
    """Zadání 12 staví další řidiče naroveň primárnímu."""
    from app.core.db import async_session_factory
    from app.models.core import User

    vehicle_id = await create_vehicle(
        logged_in_client, csrf_token, internal_code="L-11", license_plate="1LB 0011",
    )
    trip_id = await _start(logged_in_client, csrf_token, vehicle_id)

    async with async_session_factory() as db:
        other = (await db.execute(select(User).where(User.email == basic_user[0]))).scalar_one()

    detail = await logged_in_client.get(f"/kniha-jizd/trips/{trip_id}")
    added = await logged_in_client.post(
        f"/kniha-jizd/trips/{trip_id}/drivers",
        data={"csrf_token": extract_csrf_token(detail.text), "driver_id": str(other.id)},
        follow_redirects=False,
    )
    assert added.status_code == 303, added.text
    await _end(logged_in_client, csrf_token, trip_id)

    page = await logged_in_client.get(f"/kniha-jizd/logbook?driver={other.id}")
    assert _has_trip(page.text, trip_id)

    rows = _csv_lines((await logged_in_client.get(
        f"/kniha-jizd/logbook/export.csv?driver={other.id}"
    )).content)
    assert len(rows) == 2
    assert other.full_name in rows[1]


async def test_purpose_and_status_filters(logged_in_client, csrf_token):
    vehicle_id = await create_vehicle(
        logged_in_client, csrf_token, internal_code="L-12", license_plate="1LB 0012",
    )
    trip_id = await _completed_trip(logged_in_client, csrf_token, vehicle_id, purpose_code="servis")

    hit = await logged_in_client.get(f"/kniha-jizd/logbook?vehicle={vehicle_id}&purpose=servis")
    miss = await logged_in_client.get(f"/kniha-jizd/logbook?vehicle={vehicle_id}&purpose=schuzka")
    assert _has_trip(hit.text, trip_id)
    assert not _has_trip(miss.text, trip_id)

    completed = await logged_in_client.get(f"/kniha-jizd/logbook?vehicle={vehicle_id}&status=completed")
    active = await logged_in_client.get(f"/kniha-jizd/logbook?vehicle={vehicle_id}&status=active")
    assert _has_trip(completed.text, trip_id)
    assert not _has_trip(active.text, trip_id)


# --- exporty -----------------------------------------------------------

async def test_export_formats_have_the_right_headers(logged_in_client, csrf_token):
    vehicle_id = await create_vehicle(
        logged_in_client, csrf_token, internal_code="L-20", license_plate="1LB 0020",
    )
    await _completed_trip(logged_in_client, csrf_token, vehicle_id)

    expected = {
        "csv": "text/csv",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "pdf": "application/pdf",
    }
    for extension, content_type in expected.items():
        response = await logged_in_client.get(f"/kniha-jizd/logbook/export.{extension}?vehicle={vehicle_id}")
        assert response.status_code == 200, extension
        assert content_type in response.headers["content-type"], extension
        assert f'filename="kniha-jizd-{date.today():%Y-%m-%d}.{extension}"' in \
            response.headers["content-disposition"]
        # Kniha jízd je firemní údaj - nikdy do sdílené cache.
        assert "private" in response.headers["cache-control"]
        assert response.content, extension


async def test_xlsx_export_is_readable_and_filtered(logged_in_client, csrf_token):
    vehicle_id = await create_vehicle(
        logged_in_client, csrf_token, internal_code="L-21", license_plate="1LB 0021",
    )
    await _completed_trip(logged_in_client, csrf_token, vehicle_id)

    response = await logged_in_client.get(f"/kniha-jizd/logbook/export.xlsx?vehicle={vehicle_id}")
    sheet = load_workbook(io.BytesIO(response.content)).active
    headers = [cell.value for cell in sheet[1]]

    assert "Ujeté km" in headers
    assert sheet.max_row == 2
    assert sheet.cell(row=2, column=headers.index("Vozidlo") + 1).value.startswith("1LB 0021")
    assert sheet.cell(row=2, column=headers.index("Ujeté km") + 1).value == 120


async def test_pdf_export_opens_as_pdf(logged_in_client, csrf_token):
    vehicle_id = await create_vehicle(
        logged_in_client, csrf_token, internal_code="L-22", license_plate="1LB 0022",
    )
    await _completed_trip(logged_in_client, csrf_token, vehicle_id)

    response = await logged_in_client.get(f"/kniha-jizd/logbook/export.pdf?vehicle={vehicle_id}")
    assert response.content.startswith(b"%PDF-")
    assert len(response.content) > 1000


async def test_export_takes_the_whole_filter_not_just_the_page(logged_in_client, csrf_token):
    """Kdo si vyfiltruje období, čeká období - ne prvních padesát jízd."""
    from app.modules.logbook.filters import PAGE_SIZE

    vehicle_id = await create_vehicle(
        logged_in_client, csrf_token, internal_code="L-23", license_plate="1LB 0023",
    )
    count = PAGE_SIZE + 3
    for index in range(count):
        await _completed_trip(
            logged_in_client, csrf_token, vehicle_id,
            start_km=100000 + index * 10, driven=10,
        )

    page = await logged_in_client.get(f"/kniha-jizd/logbook?vehicle={vehicle_id}")
    assert "strana 1 z 2" in page.text

    rows = _csv_lines((await logged_in_client.get(
        f"/kniha-jizd/logbook/export.csv?vehicle={vehicle_id}"
    )).content)
    assert len(rows) == count + 1


# --- viditelnost (požadavek D) -----------------------------------------

async def test_hidden_vehicle_stays_out_of_list_summary_and_exports(
    logged_in_client, csrf_token, anon_client, responsible_user,
):
    """Skryté vozidlo pro cizího neexistuje - ani v XLSX.

    Kdyby se filtrovalo až v šabloně, do staženého souboru by se dostalo."""
    hidden = await create_vehicle(
        logged_in_client, csrf_token, internal_code="L-30", license_plate="1LB 0030",
        visibility="restricted",
    )
    await _completed_trip(logged_in_client, csrf_token, hidden, route_text="Tajná trasa")

    # Odpovědná osoba, která tohle vozidlo nemá na starosti.
    await login(anon_client, responsible_user)

    page = await anon_client.get("/kniha-jizd/logbook")
    assert page.status_code == 200
    assert "1LB 0030" not in page.text
    assert "Tajná trasa" not in page.text
    # Ani ve výběru vozidel - ten by prozradil, že existuje.
    assert hidden not in page.text

    for extension in ("csv", "xlsx", "pdf"):
        response = await anon_client.get(f"/kniha-jizd/logbook/export.{extension}")
        assert response.status_code == 200
        assert b"1LB 0030" not in response.content, extension

    # Ani přes přímý filtr na jeho id.
    targeted = await anon_client.get(f"/kniha-jizd/logbook/export.csv?vehicle={hidden}")
    assert len(_csv_lines(targeted.content)) == 1  # jen hlavička

    # Správce ho vidí.
    admin_page = await logged_in_client.get(f"/kniha-jizd/logbook?vehicle={hidden}")
    assert "1LB 0030" in admin_page.text


async def test_trip_of_hidden_vehicle_is_not_reachable_by_direct_link(
    logged_in_client, csrf_token, anon_client, responsible_user,
):
    """Detail jízdy dosud viditelnost vozidla nekontroloval."""
    hidden = await create_vehicle(
        logged_in_client, csrf_token, internal_code="L-31", license_plate="1LB 0031",
        visibility="restricted",
    )
    trip_id = await _completed_trip(logged_in_client, csrf_token, hidden)

    await login(anon_client, responsible_user)
    assert (await anon_client.get(f"/kniha-jizd/trips/{trip_id}")).status_code == 404

    # Správce dál 200 - kontrola nesmí zavřít vrátka i tomu, kdo právo má.
    assert (await logged_in_client.get(f"/kniha-jizd/trips/{trip_id}")).status_code == 200
